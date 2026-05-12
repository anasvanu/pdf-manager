import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import io
import os
import base64
from collections import Counter

def merge_pdfs(pdf_paths, output_path):
    merged_doc = fitz.open()
    for pdf_path in pdf_paths:
        doc = fitz.open(pdf_path)
        merged_doc.insert_pdf(doc)
        doc.close()
    merged_doc.save(output_path)
    merged_doc.close()

def _get_common_font(page):
    fonts = []
    blocks = page.get_text("dict").get("blocks", [])
    for b in blocks:
        if b['type'] == 0:
            for l in b.get("lines", []):
                for s in l.get("spans", []):
                    f = s.get("font", "").lower()
                    if "times" in f: fonts.append("Times-Roman")
                    elif "courier" in f: fonts.append("Courier")
                    else: fonts.append("Helvetica")
    if fonts:
        return Counter(fonts).most_common(1)[0][0]
    return "Helvetica"

def generate_thumbnails(pdf_path):
    doc = fitz.open(pdf_path)
    thumbnails = []
    for i in range(len(doc)):
        page = doc[i]
        pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
        img_data = pix.tobytes("png")
        b64 = base64.b64encode(img_data).decode("utf-8")
        common_font = _get_common_font(page)
        thumbnails.append({
            "id": f"page_{i}",
            "originalIndex": i,
            "data": f"data:image/png;base64,{b64}",
            "width": page.rect.width,
            "height": page.rect.height,
            "suggestedFont": common_font
        })
    doc.close()
    return thumbnails

def apply_annotations_to_page(page, annotations):
    for ann in annotations:
        type_ = ann.get('type', 'text')
        x = ann.get('x', 0)
        y = ann.get('y', 0)
        width = ann.get('width', 100)
        height = ann.get('height', 50)
        
        color_hex = ann.get('color', '#000000').lstrip('#')
        bg_hex = ann.get('bgColor', 'transparent')
        border_color_hex = ann.get('borderColor', 'transparent')
        border_width = ann.get('borderWidth', 0)
        
        color = tuple(int(color_hex[i:i+2], 16)/255.0 for i in (0, 2, 4))
        fill_color = None
        if bg_hex != 'transparent':
            fill_color = tuple(int(bg_hex.lstrip('#')[i:i+2], 16)/255.0 for i in (0, 2, 4))
            
        border_color = None
        if border_color_hex != 'transparent' and border_width > 0:
            border_color = tuple(int(border_color_hex.lstrip('#')[i:i+2], 16)/255.0 for i in (0, 2, 4))
        
        rect = fitz.Rect(x, y, x + width, y + height)

        if type_ == 'text':
            text = ann.get('text', '')
            size = ann.get('size', 14)
            is_bold = ann.get('bold', False)
            is_italic = ann.get('italic', False)
            is_underline = ann.get('underline', False)
            align_str = ann.get('align', 'left')
            font_family = ann.get('fontFamily', 'Helvetica')
            
            align_map = {'left': 0, 'center': 1, 'right': 2}
            align_val = align_map.get(align_str, 0)
            
            # Map font
            base_font = "helv"
            if font_family == "Courier": base_font = "cour"
            elif font_family == "Times-Roman": base_font = "tiro"
            
            fontname = base_font
            if is_bold and is_italic: fontname = base_font[:2] + "bi" if base_font == "helv" else base_font + "bi"
            elif is_bold: fontname = base_font[:2] + "bo" if base_font == "helv" else base_font + "bo"
            elif is_italic: fontname = base_font[:2] + "it" if base_font == "helv" else base_font + "it"
            
            if fill_color: page.draw_rect(rect, color=fill_color, fill=fill_color)
            if border_color: page.draw_rect(rect, color=border_color, width=border_width)
            
            page.insert_textbox(rect, text, fontsize=size, fontname=fontname, color=color, align=align_val)
            
            if is_underline:
                # Approximate underline position
                page.draw_line(fitz.Point(x, y + size * 1.1), fitz.Point(x + width, y + size * 1.1), color=color, width=1)
                
        elif type_ == 'rect':
            page.draw_rect(rect, color=border_color or color, fill=fill_color, width=border_width or 1)
        elif type_ == 'circle':
            page.draw_oval(rect, color=border_color or color, fill=fill_color, width=border_width or 1)
        elif type_ == 'triangle':
            p1 = fitz.Point(x + width/2, y)
            p2 = fitz.Point(x, y + height)
            p3 = fitz.Point(x + width, y + height)
            page.draw_polygon([p1, p2, p3], color=border_color or color, fill=fill_color, width=border_width or 1)
        elif type_ == 'line':
            page.draw_line(fitz.Point(x, y), fitz.Point(x + width, y + height), color=color, width=border_width or 2)
        elif type_ == 'image':
            img_b64 = ann.get('data', '')
            if ',' in img_b64: img_b64 = img_b64.split(',')[1]
            try:
                img_bytes = base64.b64decode(img_b64)
                page.insert_image(rect, stream=img_bytes)
            except Exception:
                pass

def apply_annotations_and_rearrange(pdf_path, new_order, annotations, output_path):
    doc = fitz.open(pdf_path)
    
    # Group annotations by page
    anns_by_page = {}
    if annotations:
        for ann in annotations:
            p_idx = ann.get('pageIndex')
            if p_idx not in anns_by_page: anns_by_page[p_idx] = []
            anns_by_page[p_idx].append(ann)
            
    for p_idx, anns in anns_by_page.items():
        if 0 <= p_idx < len(doc):
            apply_annotations_to_page(doc[p_idx], anns)

    new_doc = fitz.open()
    for i in new_order:
        new_doc.insert_pdf(doc, from_page=i, to_page=i)
    new_doc.save(output_path)
    new_doc.close()
    doc.close()

def _int_to_rgb(color_int):
    if isinstance(color_int, int):
        b = color_int & 255
        g = (color_int >> 8) & 255
        r = (color_int >> 16) & 255
        return RGBColor(r, g, b)
    return RGBColor(0, 0, 0)

def convert_to_pptx(pdf_path, output_path):
    doc = fitz.open(pdf_path)
    prs = Presentation()
    
    if len(doc) > 0:
        prs.slide_width = Pt(doc[0].rect.width)
        prs.slide_height = Pt(doc[0].rect.height)
    
    for page in doc:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pix = page.get_pixmap(dpi=150)
        img_stream = io.BytesIO(pix.tobytes("png"))
        slide.shapes.add_picture(img_stream, Pt(0), Pt(0), width=prs.slide_width, height=prs.slide_height)

        blocks = page.get_text("dict").get("blocks", [])
        for b in blocks:
            if b['type'] == 0: 
                x0, y0, x1, y1 = b["bbox"]
                txBox = slide.shapes.add_textbox(Pt(x0), Pt(y0), Pt(max(10, x1-x0)), Pt(max(10, y1-y0)))
                txBox.fill.solid()
                txBox.fill.fore_color.rgb = RGBColor(255, 255, 255)
                tf = txBox.text_frame
                tf.clear()
                tf.word_wrap = True
                
                text_content = []
                is_bold = is_italic = False
                size = 12
                color = RGBColor(0,0,0)
                
                for line in b["lines"]:
                    for span in line["spans"]:
                        text_content.append(span["text"])
                        font_name = span.get("font", "").lower()
                        if "bold" in font_name: is_bold = True
                        if "italic" in font_name or "oblique" in font_name: is_italic = True
                
                if b["lines"] and b["lines"][0]["spans"]:
                    size = b["lines"][0]["spans"][0]["size"]
                    if "color" in b["lines"][0]["spans"][0]:
                        color = _int_to_rgb(b["lines"][0]["spans"][0]["color"])
                
                p = tf.add_paragraph()
                run = p.add_run()
                run.text = " ".join(text_content).strip()
                run.font.size = Pt(size)
                run.font.color.rgb = color
                run.font.bold = is_bold
                run.font.italic = is_italic

    prs.save(output_path)
    doc.close()
